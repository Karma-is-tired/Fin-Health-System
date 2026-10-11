"""Financial Health & Cash-Flow API.

Settings come from environment variables (set them in your hosting dashboard):
  SECRET_KEY       long random text used to sign login tokens (required in production)
  ALLOWED_ORIGINS  comma-separated website addresses allowed to call this API
  DB_PATH          where the SQLite file lives (default database/finance_tool.db)
  CURRENCY_SYMBOL  symbol used in Word reports (default ₦)
"""
import io
import logging
import os
import re
import sqlite3
import tempfile
import time
from collections import defaultdict, deque

import pandas as pd
from docx import Document
from fastapi import Depends, FastAPI, File, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from fastapi.security import OAuth2PasswordBearer, OAuth2PasswordRequestForm
from pydantic import BaseModel
from starlette.background import BackgroundTask

import analytics
from db_setup import init_db
from security import (
    create_token, hash_password, needs_rehash, read_token, verify_password,
)

logger = logging.getLogger("finhealth")

DB_PATH = os.environ.get("DB_PATH", "database/finance_tool.db")
init_db(DB_PATH)

ALLOWED_ORIGINS = [
    origin.strip()
    for origin in os.environ.get(
        "ALLOWED_ORIGINS", "https://fin-health-system.vercel.app,http://localhost:5173"
    ).split(",")
    if origin.strip()
]
CURRENCY = os.environ.get("CURRENCY_SYMBOL", "₦")
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

app = FastAPI(title="Financial Health & Cash-Flow API")
app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["Content-Disposition"],
)


@app.exception_handler(Exception)
async def unexpected_error(request: Request, exc: Exception):
    # The real error goes to the server log, never to the visitor.
    logger.exception("Unhandled error on %s", request.url.path)
    headers = {}
    origin = request.headers.get("origin")
    if origin in ALLOWED_ORIGINS:
        headers["Access-Control-Allow-Origin"] = origin
        headers["Access-Control-Allow-Credentials"] = "true"
    return JSONResponse(
        status_code=500,
        content={"detail": "Something went wrong on our side. Please try again."},
        headers=headers,
    )


def get_db():
    conn = sqlite3.connect(DB_PATH, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    try:
        yield conn
    finally:
        conn.close()


oauth2_scheme = OAuth2PasswordBearer(tokenUrl="login")
SIGN_IN_AGAIN = "Your session has expired. Please sign in again."


def get_current_user(token: str = Depends(oauth2_scheme), db: sqlite3.Connection = Depends(get_db)):
    email = read_token(token)
    if email is None:
        raise HTTPException(401, SIGN_IN_AGAIN, headers={"WWW-Authenticate": "Bearer"})
    user = db.execute("SELECT * FROM users WHERE lower(email) = ?", (email.lower(),)).fetchone()
    if user is None:
        raise HTTPException(401, SIGN_IN_AGAIN, headers={"WWW-Authenticate": "Bearer"})
    return user


# ---------------------------------------------------------------------------
# Accounts: email + password, nothing else
# ---------------------------------------------------------------------------
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
MAX_FAILED_LOGINS = 5
LOCKOUT_SECONDS = 15 * 60
_failed_logins = defaultdict(deque)
_DUMMY_HASH = hash_password("not-a-real-password")  # keeps timing equal for unknown emails


class Credentials(BaseModel):
    email: str
    password: str


def _recent_failures(key):
    now = time.time()
    attempts = _failed_logins[key]
    while attempts and now - attempts[0] > LOCKOUT_SECONDS:
        attempts.popleft()
    return attempts


@app.get("/")
def root():
    return {"message": "Financial Health API is running."}


@app.post("/signup")
def signup(body: Credentials, db: sqlite3.Connection = Depends(get_db)):
    email = body.email.strip().lower()
    if len(email) > 254 or not EMAIL_RE.match(email):
        raise HTTPException(400, "Please enter a valid email address.")
    if not 8 <= len(body.password) <= 128:
        raise HTTPException(400, "Your password must be between 8 and 128 characters.")

    taken = "That email is already registered. Try signing in instead."
    if db.execute("SELECT 1 FROM users WHERE lower(email) = ?", (email,)).fetchone():
        raise HTTPException(400, taken)
    try:
        db.execute(
            "INSERT INTO users (email, password_hash, is_verified) VALUES (?, ?, 1)",
            (email, hash_password(body.password)),
        )
        db.commit()
    except sqlite3.IntegrityError:
        raise HTTPException(400, taken)

    # Signing up also signs you in, so you land straight on your dashboard.
    return {"message": "Account created.", "access_token": create_token(email), "token_type": "bearer"}


@app.post("/login")
def login(form: OAuth2PasswordRequestForm = Depends(), db: sqlite3.Connection = Depends(get_db)):
    email = form.username.strip().lower()
    if len(_recent_failures(email)) >= MAX_FAILED_LOGINS:
        raise HTTPException(429, "Too many failed attempts. Please wait 15 minutes and try again.")

    user = db.execute("SELECT * FROM users WHERE lower(email) = ?", (email,)).fetchone()
    valid = verify_password(form.password, user["password_hash"] if user else _DUMMY_HASH)
    if not user or not valid:
        _recent_failures(email).append(time.time())
        raise HTTPException(401, "Incorrect email or password.")

    _failed_logins.pop(email, None)
    if needs_rehash(user["password_hash"]):  # upgrade old accounts on their next login
        db.execute("UPDATE users SET password_hash = ? WHERE id = ?", (hash_password(form.password), user["id"]))
        db.commit()
    return {"access_token": create_token(user["email"]), "token_type": "bearer"}


@app.get("/me")
def me(user=Depends(get_current_user)):
    return {"email": user["email"]}


# ---------------------------------------------------------------------------
# Analytics (every query is limited to the logged-in user's own rows)
# ---------------------------------------------------------------------------
@app.get("/health-score")
def health_score(user=Depends(get_current_user), db: sqlite3.Connection = Depends(get_db)):
    return analytics.compute_scores(db, user["id"])


@app.get("/financial-summary")
def financial_summary(user=Depends(get_current_user), db: sqlite3.Connection = Depends(get_db)):
    return analytics.financial_summary(db, user["id"])


@app.get("/monthly-trend")
def monthly_trend(user=Depends(get_current_user), db: sqlite3.Connection = Depends(get_db)):
    return analytics.monthly_trend(db, user["id"])


def _read_table(filename: str, contents: bytes) -> pd.DataFrame:
    name = (filename or "").lower()
    try:
        if name.endswith(".csv"):
            try:
                return pd.read_csv(io.BytesIO(contents), encoding="utf-8-sig")
            except UnicodeDecodeError:
                return pd.read_csv(io.BytesIO(contents), encoding="latin-1")
        if name.endswith(".xlsx"):
            return pd.read_excel(io.BytesIO(contents))
    except Exception:
        logger.exception("Could not read uploaded file")
        raise HTTPException(400, "We couldn't read that file. Make sure it is a valid CSV or Excel file.")
    raise HTTPException(400, "Please upload a .csv or .xlsx file.")


@app.post("/upload/{kind}")
def upload(
    kind: str,
    file: UploadFile = File(...),
    user=Depends(get_current_user),
    db: sqlite3.Connection = Depends(get_db),
):
    if kind not in analytics.NORMALIZERS:
        raise HTTPException(404, "Unknown upload type.")
    contents = file.file.read(MAX_UPLOAD_BYTES + 1)
    if len(contents) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "That file is too large. The limit is 5 MB.")

    raw = _read_table(file.filename, contents)
    try:
        cleaned, skipped = analytics.NORMALIZERS[kind](raw)
    except ValueError as err:
        raise HTTPException(400, str(err))

    saved = analytics.save_upload(db, user["id"], kind, file.filename, cleaned)
    message = f"Added {saved} {kind} row{'s' if saved != 1 else ''}."
    if skipped:
        message += f" Skipped {skipped} blank or unreadable row{'s' if skipped != 1 else ''}."
    return {"message": message}


@app.post("/load-demo-data")
def load_demo_data(user=Depends(get_current_user), db: sqlite3.Connection = Depends(get_db)):
    count = analytics.load_demo_data(db, user["id"])
    return {"message": f"Loaded {count} sample transactions, invoices and a budget."}


# ---------------------------------------------------------------------------
# Word reports + history
# ---------------------------------------------------------------------------
def _money(value):
    return f"{CURRENCY}{float(value or 0):,.2f}"


def _delete_quietly(path):
    try:
        os.remove(path)
    except OSError:
        pass


def _docx_response(month, overall, revenue, profit, filename):
    doc = Document()
    doc.add_heading("Executive Financial Health Report", 0)
    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(f"Analysis period: {month}")
    line = doc.add_paragraph("Overall Financial Health Score: ")
    line.add_run(f"{overall}/100").bold = True
    line.add_run(f" ({analytics.band(overall)})")

    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Metric", "Value"
    for label, value in (("Revenue", revenue), ("Net profit", profit)):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, _money(value)

    doc.add_heading("2. Automated Strategic Advice", level=1)
    if profit < revenue * 0.1:
        doc.add_paragraph("PROFITABILITY: Margins are tight. Review non-essential spending soon.")
    else:
        doc.add_paragraph("PROFITABILITY: Margins are healthy. Consider reinvesting surplus cash.")

    fd, path = tempfile.mkstemp(suffix=".docx")
    with os.fdopen(fd, "wb") as handle:
        doc.save(handle)
    # The temporary file is removed as soon as it has been sent.
    return FileResponse(path, filename=filename, media_type=DOCX_TYPE, background=BackgroundTask(_delete_quietly, path))


@app.get("/download-report")
def download_report(user=Depends(get_current_user), db: sqlite3.Connection = Depends(get_db)):
    summary = analytics.financial_summary(db, user["id"])
    if not summary["has_data"]:
        raise HTTPException(400, "Upload some transactions before generating a report.")
    overall = analytics.compute_scores(db, user["id"])["overall"]

    db.execute(
        "INSERT INTO report_history (user_id, report_month, overall_score, revenue, net_profit) VALUES (?, ?, ?, ?, ?)",
        (user["id"], summary["latest_month"], overall, summary["revenue"], summary["net_profit"]),
    )
    db.commit()
    return _docx_response(
        summary["latest_month"], overall, summary["revenue"], summary["net_profit"],
        f"FinHealth_{summary['latest_month']}.docx",
    )


@app.get("/history")
def history(user=Depends(get_current_user), db: sqlite3.Connection = Depends(get_db)):
    rows = db.execute(
        "SELECT id, report_month, overall_score, revenue, net_profit, created_at "
        "FROM report_history WHERE user_id = ? ORDER BY created_at DESC, id DESC",
        (user["id"],),
    ).fetchall()
    return [dict(row) for row in rows]


@app.get("/download-past-report/{report_id}")
def download_past_report(report_id: int, user=Depends(get_current_user), db: sqlite3.Connection = Depends(get_db)):
    report = db.execute(
        "SELECT * FROM report_history WHERE id = ? AND user_id = ?", (report_id, user["id"])
    ).fetchone()
    if not report:
        raise HTTPException(404, "Report not found.")
    return _docx_response(
        report["report_month"], report["overall_score"], report["revenue"], report["net_profit"],
        f"FinHealth_Archive_{report['report_month']}.docx",
    )
